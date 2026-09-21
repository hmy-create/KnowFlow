from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]

PATH = (
    ROOT
    / "eval"
    / "frozen"
    / "KnowFlow_Eval_Dataset_v1.0.xlsx"
)


def main():

    if not PATH.exists():
        raise FileNotFoundError(PATH)

    workbook = load_workbook(
        PATH,
        read_only=True,
        data_only=True,
    )

    print("=" * 100)
    print("KnowFlow Frozen Eval Dataset")
    print("=" * 100)

    print(f"path={PATH}")
    print(f"sheets={workbook.sheetnames}")

    for sheet_name in workbook.sheetnames:

        ws = workbook[sheet_name]

        rows = list(
            ws.iter_rows(
                values_only=True
            )
        )

        non_empty_rows = [
            row
            for row in rows
            if any(
                value is not None
                for value in row
            )
        ]

        print()
        print("-" * 100)
        print(f"SHEET: {sheet_name}")
        print(
            f"non_empty_rows="
            f"{len(non_empty_rows)}"
        )

        if not non_empty_rows:
            continue

        print()
        print("HEADER:")
        print(non_empty_rows[0])

        print()
        print("FIRST 5 DATA ROWS:")

        for row in non_empty_rows[1:6]:
            print(row)


if __name__ == "__main__":
    main()