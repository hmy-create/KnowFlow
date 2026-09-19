from pathlib import Path

from docx import Document


REPO_ROOT = Path(__file__).resolve().parents[1]

SOURCE_FILE = (
    REPO_ROOT
    / "docs"
    / "corpus"
    / "KnowFlow_Corpus_V1.0"
    / "00_语料说明与测试用例.docx"
)

OUTPUT_FILE = (
    REPO_ROOT
    / "docs"
    / "corpus"
    / "corpus_guide_dump.txt"
)


def clean_text(text: str) -> str:
    return " ".join(
        text.replace("\n", " ").split()
    )


def main():
    if not SOURCE_FILE.exists():
        raise FileNotFoundError(
            f"找不到语料说明文件：{SOURCE_FILE}"
        )

    doc = Document(SOURCE_FILE)

    lines = []

    lines.append(
        "===== KnowFlow Corpus Guide ====="
    )
    lines.append(
        f"source: {SOURCE_FILE.name}"
    )
    lines.append("")

    # --------------------------
    # 普通段落
    # --------------------------
    lines.append(
        "===== PARAGRAPHS ====="
    )

    for index, paragraph in enumerate(
        doc.paragraphs,
        start=1,
    ):
        text = clean_text(
            paragraph.text
        )

        if not text:
            continue

        style_name = (
            paragraph.style.name
            if paragraph.style
            else ""
        )

        lines.append(
            f"[P{index}] "
            f"[{style_name}] "
            f"{text}"
        )

    lines.append("")

    # --------------------------
    # 表格
    # --------------------------
    lines.append(
        "===== TABLES ====="
    )

    for table_index, table in enumerate(
        doc.tables,
        start=1,
    ):
        lines.append("")
        lines.append(
            f"--- TABLE {table_index} ---"
        )

        for row_index, row in enumerate(
            table.rows,
            start=1,
        ):
            cells = [
                clean_text(cell.text)
                for cell in row.cells
            ]

            lines.append(
                f"[R{row_index}] "
                + " | ".join(cells)
            )

    OUTPUT_FILE.write_text(
        "\n".join(lines),
        encoding="utf-8",
    )

    print(
        "Corpus guide extracted:"
    )
    print(
        OUTPUT_FILE
    )


if __name__ == "__main__":
    main()