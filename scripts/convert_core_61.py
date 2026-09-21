from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


# ================================================================
# Paths
# ================================================================
REPO_ROOT = Path(
    __file__
).resolve().parents[1]

SOURCE_XLSX = (
    REPO_ROOT
    / "eval"
    / "frozen"
    / "KnowFlow_Eval_Dataset_v1.0.xlsx"
)

OUTPUT_CSV = (
    REPO_ROOT
    / "eval"
    / "core_61.csv"
)

REPORT_JSON = (
    REPO_ROOT
    / "eval"
    / "core_61_conversion_report.json"
)

SOURCE_SHEET = "Eval_Cases"

EXPECTED_CASE_COUNT = 61


# ================================================================
# Canonical Output Schema
#
# 前半部分：
# 后续 S9 自动 Eval 直接使用。
#
# 后半部分 frozen_*：
# 保留 Dify RUN_004 冻结记录，
# 避免转换时丢失原始基线信息。
# ================================================================
OUTPUT_FIELDS = [
    "case_id",
    "set_name",
    "category",
    "query",

    "user_id",
    "department",
    "role",
    "authorized_kb_ids_json",

    "expected_domain",
    "expected_sensitivity",
    "expected_version_mode",
    "expected_permission_decision",
    "expected_decision",

    "relevant_document_ids_json",
    "relevant_chunk_ids_json",

    "expected_citation_document_ids_json",
    "expected_citation_chunk_ids_json",

    "require_clarifying_question",

    "expected_answer_keypoints",
    "expected_source",
    "route_expectation",

    "notes",

    # ------------------------------------------------------------
    # Frozen Dify Baseline Audit Fields
    # ------------------------------------------------------------
    "frozen_stage",
    "frozen_domain_actual",
    "frozen_actual_decision",
    "frozen_actual_result_summary",
    "frozen_actual_source",
    "frozen_route_actual",
    "frozen_pass",
    "test_type",
]


# ================================================================
# Expected Excel Headers
# ================================================================
REQUIRED_SOURCE_HEADERS = {
    "case_id",
    "stage",
    "domain_expected",
    "domain_actual",
    "user_role",
    "version_expected",
    "sensitivity_expected",
    "query",
    "expected_decision",
    "actual_decision",
    "expected_answer_keypoints",
    "actual_result_summary",
    "expected_source",
    "actual_source",
    "route_expectation",
    "route_actual",
    "pass",
    "notes",
    "test_type",
}


# ================================================================
# Domain Normalization
# ================================================================
DOMAIN_MAP = {
    "finance": "Finance",
    "hr": "HR",
    "product": "Product",
    "service": "Service",
    "other": "Other",
}


# ================================================================
# Version Normalization
#
# Excel:
#   current
#   historical
#   comparison
#
# FastAPI:
#   Current
#   Historical
#   Comparison
# ================================================================
VERSION_MAP = {
    "current": "Current",
    "historical": "Historical",
    "comparison": "Comparison",

    # 允许冻结数据显式写 none / null
    "none": "",
    "null": "",
    "": "",
}


# ================================================================
# Sensitivity Normalization
# ================================================================
SENSITIVITY_MAP = {
    "normal": "Normal",
    "restricted": "Restricted",

    "none": "",
    "null": "",
    "": "",
}


# ================================================================
# Decision Normalization
# ================================================================
DECISION_MAP = {
    "answer": "answer",
    "clarify": "clarify",
    "conflict": "conflict",
    "refuse": "refuse",
    "no_access": "no_access",
    "no access": "no_access",
    "no-access": "no_access",
}


# ================================================================
# Evaluation Runtime KB Fixtures
#
# 注意：
# 这些不是 Excel Gold。
#
# Excel 只提供 user_role，
# 但后续 /chat 需要完整 user：
#
#   user_id
#   department
#   role
#   authorized_kb_ids
#
# 因此这里仅构造测试执行身份。
# ================================================================
PUBLIC_KB_BY_DOMAIN = {
    "Finance": "KB_FINANCE_PUBLIC",
    "HR": "KB_HR_PUBLIC",
    "Product": "KB_PRODUCT_PUBLIC",
    "Service": "KB_SERVICE_PUBLIC",
}


# ================================================================
# Helpers
# ================================================================
def clean_text(
    value: Any,
) -> str:
    if value is None:
        return ""

    return str(
        value
    ).strip()


def normalize_key(
    value: Any,
) -> str:
    return clean_text(
        value
    ).lower()


def json_list(
    values: list[str],
) -> str:
    return json.dumps(
        values,
        ensure_ascii=False,
        separators=(",", ":"),
    )


def normalize_domain(
    value: Any,
) -> str:

    raw = normalize_key(
        value
    )

    if not raw:
        return ""

    if raw not in DOMAIN_MAP:
        raise ValueError(
            "Unknown domain_expected: "
            f"{value!r}"
        )

    return DOMAIN_MAP[
        raw
    ]


def normalize_version(
    value: Any,
) -> str:

    raw = normalize_key(
        value
    )

    if raw not in VERSION_MAP:
        raise ValueError(
            "Unknown version_expected: "
            f"{value!r}"
        )

    return VERSION_MAP[
        raw
    ]


def normalize_sensitivity(
    value: Any,
) -> str:

    raw = normalize_key(
        value
    )

    if raw not in SENSITIVITY_MAP:
        raise ValueError(
            "Unknown sensitivity_expected: "
            f"{value!r}"
        )

    return SENSITIVITY_MAP[
        raw
    ]


def normalize_decision(
    value: Any,
) -> str:

    raw = normalize_key(
        value
    )

    if not raw:
        raise ValueError(
            "expected_decision is empty"
        )

    if raw not in DECISION_MAP:
        raise ValueError(
            "Unknown expected_decision: "
            f"{value!r}"
        )

    return DECISION_MAP[
        raw
    ]


def normalize_bool(
    value: Any,
) -> bool:

    if isinstance(
        value,
        bool,
    ):
        return value

    if isinstance(
        value,
        (
            int,
            float,
        ),
    ):
        return bool(
            value
        )

    raw = normalize_key(
        value
    )

    if raw in {
        "true",
        "1",
        "yes",
        "y",
        "pass",
        "passed",
    }:
        return True

    if raw in {
        "false",
        "0",
        "no",
        "n",
        "fail",
        "failed",
    }:
        return False

    raise ValueError(
        "Cannot normalize boolean: "
        f"{value!r}"
    )


# ================================================================
# Category
#
# 不重新定义冻结 Case。
#
# 优先使用 Case ID 前缀，
# 只是给后续报告分组使用。
#
# 如果无法可靠识别，则写 "baseline"，
# 不影响 Gold。
# ================================================================
def infer_category(
    case_id: str,
    expected_decision: str,
    version_mode: str,
    sensitivity: str,
) -> str:

    upper_id = case_id.upper()

    # 明确的版本边界 / version case
    if upper_id.startswith("V"):
        return "version"

    # Comparison
    if (
        version_mode
        == "Comparison"
    ):
        return "comparison"

    # Permission / restricted
    if (
        expected_decision
        == "no_access"
        or sensitivity
        == "Restricted"
    ):
        return "permission"

    # Clarify
    if (
        expected_decision
        == "clarify"
    ):
        return "clarify"

    # Conflict
    if (
        expected_decision
        == "conflict"
    ):
        return "conflict"

    # Refuse
    if (
        expected_decision
        == "refuse"
    ):
        return "refuse"

    # Historical
    if (
        version_mode
        == "Historical"
    ):
        return "historical"

    # Current
    if (
        version_mode
        == "Current"
    ):
        return "current"

    return "baseline"


# ================================================================
# Runtime User Fixture
#
# 这里仅用于让后续 Eval 能真正调用 /chat。
#
# 不参与 Frozen Gold Accuracy。
# ================================================================
def build_runtime_user(
    case_id: str,
    domain: str,
    role: str,
) -> dict:

    normalized_role = (
        role.strip()
        or "employee"
    )

    user_id = (
        f"EVAL_{case_id}"
    )

    if domain == "Other":
        department = "general"
        authorized_kb_ids = []

        return {
            "user_id":
                user_id,
            "department":
                department,
            "role":
                normalized_role,
            "authorized_kb_ids":
                authorized_kb_ids,
        }

    department_map = {
        "Finance": "finance",
        "HR": "hr",
        "Product": "product",
        "Service": "service",
    }

    department = (
        department_map.get(
            domain,
            "general",
        )
    )

    authorized_kb_ids = []

    public_kb = (
        PUBLIC_KB_BY_DOMAIN.get(
            domain
        )
    )

    if public_kb:
        authorized_kb_ids.append(
            public_kb
        )

    role_key = (
        normalized_role
        .lower()
        .replace("-", "_")
        .replace(" ", "_")
    )

    # Finance restricted corpus
    #
    # 当前 KnowFlow 冻结语料中，
    # Private KB 是 Finance Restricted。
    if (
        domain == "Finance"
        and role_key
        in {
            "finance_manager",
            "admin",
        }
    ):
        authorized_kb_ids.append(
            "KB_FINANCE_PRIVATE"
        )

    # Admin 作为执行 fixture：
    # 若数据集中存在 admin，
    # 给它已知的全部公共知识库。
    if role_key == "admin":

        authorized_kb_ids = [
            "KB_FINANCE_PUBLIC",
            "KB_HR_PUBLIC",
            "KB_PRODUCT_PUBLIC",
            "KB_SERVICE_PUBLIC",
            "KB_FINANCE_PRIVATE",
        ]

    return {
        "user_id":
            user_id,
        "department":
            department,
        "role":
            normalized_role,
        "authorized_kb_ids":
            authorized_kb_ids,
    }


# ================================================================
# Excel Reader
# ================================================================
def load_eval_cases() -> list[dict]:

    if not SOURCE_XLSX.exists():
        raise FileNotFoundError(
            SOURCE_XLSX
        )

    workbook = load_workbook(
        SOURCE_XLSX,
        read_only=True,
        data_only=True,
    )

    if (
        SOURCE_SHEET
        not in workbook.sheetnames
    ):
        raise RuntimeError(
            f"Sheet {SOURCE_SHEET!r} "
            "not found. "
            f"Available sheets: "
            f"{workbook.sheetnames}"
        )

    ws = workbook[
        SOURCE_SHEET
    ]

    rows = list(
        ws.iter_rows(
            values_only=True
        )
    )

    non_empty_rows = [
        row
        for row in rows
        if any(
            value is not None
            for value in row
        )
    ]

    if not non_empty_rows:
        raise RuntimeError(
            "Eval_Cases sheet is empty."
        )

    headers = [
        clean_text(value)
        for value
        in non_empty_rows[0]
    ]

    header_set = set(
        headers
    )

    missing_headers = (
        REQUIRED_SOURCE_HEADERS
        - header_set
    )

    if missing_headers:
        raise RuntimeError(
            "Frozen Eval sheet is missing "
            "required headers: "
            f"{sorted(missing_headers)}"
        )

    cases = []

    for row_number, row in enumerate(
        non_empty_rows[1:],
        start=2,
    ):

        padded = list(
            row
        )

        if (
            len(padded)
            < len(headers)
        ):
            padded.extend(
                [None]
                * (
                    len(headers)
                    - len(padded)
                )
            )

        record = {
            header: padded[index]
            for index, header
            in enumerate(headers)
        }

        record[
            "__row_number__"
        ] = row_number

        cases.append(
            record
        )

    return cases


# ================================================================
# Frozen Dataset Validation
# ================================================================
def validate_source_cases(
    cases: list[dict],
) -> None:

    errors = []

    if (
        len(cases)
        != EXPECTED_CASE_COUNT
    ):
        errors.append(
            "Expected exactly "
            f"{EXPECTED_CASE_COUNT} cases, "
            f"but found {len(cases)}."
        )

    case_ids = []

    for record in cases:

        row_number = record[
            "__row_number__"
        ]

        case_id = clean_text(
            record.get(
                "case_id"
            )
        )

        query = clean_text(
            record.get(
                "query"
            )
        )

        domain = clean_text(
            record.get(
                "domain_expected"
            )
        )

        decision = clean_text(
            record.get(
                "expected_decision"
            )
        )

        role = clean_text(
            record.get(
                "user_role"
            )
        )

        if not case_id:
            errors.append(
                f"Row {row_number}: "
                "case_id is empty."
            )
        else:
            case_ids.append(
                case_id
            )

        if not query:
            errors.append(
                f"Row {row_number} "
                f"({case_id}): "
                "query is empty."
            )

        # Frozen Gold 中真正不能缺失的字段
        if not domain:
            errors.append(
                f"Row {row_number} "
                f"({case_id}): "
                "domain_expected is empty."
            )

        if not decision:
            errors.append(
                f"Row {row_number} "
                f"({case_id}): "
                "expected_decision is empty."
            )

        if not role:
            errors.append(
                f"Row {row_number} "
                f"({case_id}): "
                "user_role is empty."
            )

        # 验证枚举是否合法。
        #
        # version / sensitivity 允许空值，
        # 但若非空必须属于已知枚举。
        try:
            normalize_domain(
                domain
            )
        except ValueError as exc:
            errors.append(
                f"Row {row_number} "
                f"({case_id}): {exc}"
            )

        try:
            normalize_version(
                record.get(
                    "version_expected"
                )
            )
        except ValueError as exc:
            errors.append(
                f"Row {row_number} "
                f"({case_id}): {exc}"
            )

        try:
            normalize_sensitivity(
                record.get(
                    "sensitivity_expected"
                )
            )
        except ValueError as exc:
            errors.append(
                f"Row {row_number} "
                f"({case_id}): {exc}"
            )

        try:
            normalize_decision(
                decision
            )
        except ValueError as exc:
            errors.append(
                f"Row {row_number} "
                f"({case_id}): {exc}"
            )

        # RUN_004 FINAL 应该是封版通过集。
        try:
            frozen_pass = normalize_bool(
                record.get(
                    "pass"
                )
            )

            if not frozen_pass:
                errors.append(
                    f"Row {row_number} "
                    f"({case_id}): "
                    "frozen pass is False. "
                    "This conflicts with "
                    "RUN_004 FINAL 61/61."
                )

        except ValueError as exc:
            errors.append(
                f"Row {row_number} "
                f"({case_id}): {exc}"
            )

    duplicates = [
        case_id
        for case_id, count
        in Counter(
            case_ids
        ).items()
        if count > 1
    ]

    if duplicates:
        errors.append(
            "Duplicate case_id values: "
            f"{sorted(duplicates)}"
        )

    if errors:

        message = (
            "\n".join(
                [
                    "=" * 80,
                    "Frozen Eval validation FAILED",
                    "=" * 80,
                    *[
                        f"- {error}"
                        for error in errors
                    ],
                ]
            )
        )

        raise RuntimeError(
            message
        )


# ================================================================
# Conversion
# ================================================================
def convert_case(
    source: dict,
) -> dict:

    case_id = clean_text(
        source[
            "case_id"
        ]
    )

    domain = normalize_domain(
        source[
            "domain_expected"
        ]
    )

    version_mode = (
        normalize_version(
            source.get(
                "version_expected"
            )
        )
    )

    sensitivity = (
        normalize_sensitivity(
            source.get(
                "sensitivity_expected"
            )
        )
    )

    expected_decision = (
        normalize_decision(
            source[
                "expected_decision"
            ]
        )
    )

    role = clean_text(
        source.get(
            "user_role"
        )
    )

    runtime_user = (
        build_runtime_user(
            case_id=case_id,
            domain=domain,
            role=role,
        )
    )

    category = infer_category(
        case_id=case_id,
        expected_decision=(
            expected_decision
        ),
        version_mode=(
            version_mode
        ),
        sensitivity=(
            sensitivity
        ),
    )

    require_clarify = (
        expected_decision
        == "clarify"
    )

    frozen_pass = normalize_bool(
        source.get(
            "pass"
        )
    )

    # ============================================================
    # IMPORTANT
    #
    # 下列 Gold 在当前 Excel 中没有结构化 ID：
    #
    # relevant_document_ids
    # relevant_chunk_ids
    # expected_citation_document_ids
    # expected_citation_chunk_ids
    #
    # 虽然 expected_source 有文件名，
    # 但这里不擅自把文件名映射成 document_id，
    # 更不从当前 FastAPI 检索结果倒推 Gold。
    #
    # 因此安全地保持 []。
    # ============================================================
    return {
        "case_id":
            case_id,

        "set_name":
            "core_61",

        "category":
            category,

        "query":
            clean_text(
                source.get(
                    "query"
                )
            ),

        "user_id":
            runtime_user[
                "user_id"
            ],

        "department":
            runtime_user[
                "department"
            ],

        "role":
            runtime_user[
                "role"
            ],

        "authorized_kb_ids_json":
            json_list(
                runtime_user[
                    "authorized_kb_ids"
                ]
            ),

        "expected_domain":
            domain,

        "expected_sensitivity":
            sensitivity,

        "expected_version_mode":
            version_mode,

        # Excel 没有独立的
        # permission_decision Gold 列。
        #
        # 不从 expected_decision
        # 或当前 FastAPI 结果倒推。
        "expected_permission_decision":
            "",

        "expected_decision":
            expected_decision,

        "relevant_document_ids_json":
            json_list([]),

        "relevant_chunk_ids_json":
            json_list([]),

        "expected_citation_document_ids_json":
            json_list([]),

        "expected_citation_chunk_ids_json":
            json_list([]),

        "require_clarifying_question":
            (
                "true"
                if require_clarify
                else "false"
            ),

        "expected_answer_keypoints":
            clean_text(
                source.get(
                    "expected_answer_keypoints"
                )
            ),

        "expected_source":
            clean_text(
                source.get(
                    "expected_source"
                )
            ),

        "route_expectation":
            clean_text(
                source.get(
                    "route_expectation"
                )
            ),

        "notes":
            clean_text(
                source.get(
                    "notes"
                )
            ),

        "frozen_stage":
            clean_text(
                source.get(
                    "stage"
                )
            ),

        "frozen_domain_actual":
            clean_text(
                source.get(
                    "domain_actual"
                )
            ),

        "frozen_actual_decision":
            clean_text(
                source.get(
                    "actual_decision"
                )
            ),

        "frozen_actual_result_summary":
            clean_text(
                source.get(
                    "actual_result_summary"
                )
            ),

        "frozen_actual_source":
            clean_text(
                source.get(
                    "actual_source"
                )
            ),

        "frozen_route_actual":
            clean_text(
                source.get(
                    "route_actual"
                )
            ),

        "frozen_pass":
            (
                "true"
                if frozen_pass
                else "false"
            ),

        "test_type":
            clean_text(
                source.get(
                    "test_type"
                )
            ),
    }


# ================================================================
# Converted Dataset Validation
# ================================================================
def validate_converted_cases(
    cases: list[dict],
) -> dict:

    errors = []

    case_ids = [
        item["case_id"]
        for item in cases
    ]

    if (
        len(cases)
        != EXPECTED_CASE_COUNT
    ):
        errors.append(
            f"Converted case count is "
            f"{len(cases)}, expected "
            f"{EXPECTED_CASE_COUNT}."
        )

    if (
        len(set(case_ids))
        != len(case_ids)
    ):
        errors.append(
            "Converted case IDs "
            "are not unique."
        )

    empty_queries = [
        item["case_id"]
        for item in cases
        if not item["query"]
    ]

    if empty_queries:
        errors.append(
            "Empty query cases: "
            f"{empty_queries}"
        )

    missing_domains = [
        item["case_id"]
        for item in cases
        if not item[
            "expected_domain"
        ]
    ]

    if missing_domains:
        errors.append(
            "Missing expected_domain: "
            f"{missing_domains}"
        )

    missing_decisions = [
        item["case_id"]
        for item in cases
        if not item[
            "expected_decision"
        ]
    ]

    if missing_decisions:
        errors.append(
            "Missing expected_decision: "
            f"{missing_decisions}"
        )

    frozen_failed = [
        item["case_id"]
        for item in cases
        if item[
            "frozen_pass"
        ] != "true"
    ]

    if frozen_failed:
        errors.append(
            "Frozen RUN_004 failed cases: "
            f"{frozen_failed}"
        )

    if errors:
        raise RuntimeError(
            "\n".join(
                [
                    "=" * 80,
                    "Converted Core 61 validation FAILED",
                    "=" * 80,
                    *[
                        f"- {error}"
                        for error in errors
                    ],
                ]
            )
        )

    version_counts = Counter(
        item[
            "expected_version_mode"
        ]
        or "<blank>"
        for item in cases
    )

    sensitivity_counts = Counter(
        item[
            "expected_sensitivity"
        ]
        or "<blank>"
        for item in cases
    )

    decision_counts = Counter(
        item[
            "expected_decision"
        ]
        for item in cases
    )

    domain_counts = Counter(
        item[
            "expected_domain"
        ]
        for item in cases
    )

    role_counts = Counter(
        item[
            "role"
        ]
        for item in cases
    )

    category_counts = Counter(
        item[
            "category"
        ]
        for item in cases
    )

    clarify_cases = [
        item["case_id"]
        for item in cases
        if (
            item[
                "require_clarifying_question"
            ]
            == "true"
        )
    ]

    return {
        "source_file":
            str(SOURCE_XLSX),

        "source_sheet":
            SOURCE_SHEET,

        "output_file":
            str(OUTPUT_CSV),

        "expected_case_count":
            EXPECTED_CASE_COUNT,

        "actual_case_count":
            len(cases),

        "case_ids_unique":
            (
                len(set(case_ids))
                == len(case_ids)
            ),

        "empty_query_count":
            len(empty_queries),

        "missing_expected_domain_count":
            len(missing_domains),

        "missing_expected_decision_count":
            len(missing_decisions),

        "frozen_failed_count":
            len(frozen_failed),

        "domain_counts":
            dict(domain_counts),

        "version_counts":
            dict(version_counts),

        "sensitivity_counts":
            dict(
                sensitivity_counts
            ),

        "decision_counts":
            dict(decision_counts),

        "role_counts":
            dict(role_counts),

        "category_counts":
            dict(category_counts),

        "clarify_case_count":
            len(clarify_cases),

        "clarify_case_ids":
            clarify_cases,

        # 明确记录哪些 Gold
        # 当前没有从 Excel 获得。
        "unavailable_structured_gold": [
            "expected_permission_decision",
            "relevant_document_ids",
            "relevant_chunk_ids",
            "expected_citation_document_ids",
            "expected_citation_chunk_ids",
        ],

        "conversion_policy": {
            "gold_source":
                "Frozen Eval_Cases sheet",

            "runtime_user_fixture":
                (
                    "Generated only for "
                    "executing FastAPI /chat; "
                    "not treated as Gold."
                ),

            "missing_gold_policy":
                (
                    "Do not infer from "
                    "current FastAPI output."
                ),

            "expected_source_policy":
                (
                    "Original filename-level "
                    "expected_source is preserved "
                    "verbatim; it is not silently "
                    "converted to document_id."
                ),
        },

        "validation_passed":
            True,
    }


# ================================================================
# Writers
# ================================================================
def write_csv(
    cases: list[dict],
) -> None:

    OUTPUT_CSV.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with OUTPUT_CSV.open(
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as file:

        writer = csv.DictWriter(
            file,
            fieldnames=(
                OUTPUT_FIELDS
            ),
            extrasaction="raise",
        )

        writer.writeheader()

        writer.writerows(
            cases
        )


def write_report(
    report: dict,
) -> None:

    REPORT_JSON.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    REPORT_JSON.write_text(
        json.dumps(
            report,
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )


# ================================================================
# Console Summary
# ================================================================
def print_summary(
    report: dict,
) -> None:

    print()
    print("=" * 100)
    print(
        "KnowFlow Core 61 Conversion"
    )
    print("=" * 100)

    print(
        f"source : "
        f"{report['source_file']}"
    )

    print(
        f"sheet  : "
        f"{report['source_sheet']}"
    )

    print(
        f"output : "
        f"{report['output_file']}"
    )

    print()

    print(
        "cases             : "
        f"{report['actual_case_count']}"
    )

    print(
        "case_ids_unique   : "
        f"{report['case_ids_unique']}"
    )

    print(
        "empty_queries     : "
        f"{report['empty_query_count']}"
    )

    print(
        "missing_domain    : "
        f"{report['missing_expected_domain_count']}"
    )

    print(
        "missing_decision  : "
        f"{report['missing_expected_decision_count']}"
    )

    print(
        "frozen_failed     : "
        f"{report['frozen_failed_count']}"
    )

    print()

    print(
        "domain_counts:"
    )

    for key, value in (
        report[
            "domain_counts"
        ].items()
    ):
        print(
            f"  {key}: {value}"
        )

    print()

    print(
        "version_counts:"
    )

    for key, value in (
        report[
            "version_counts"
        ].items()
    ):
        print(
            f"  {key}: {value}"
        )

    print()

    print(
        "sensitivity_counts:"
    )

    for key, value in (
        report[
            "sensitivity_counts"
        ].items()
    ):
        print(
            f"  {key}: {value}"
        )

    print()

    print(
        "decision_counts:"
    )

    for key, value in (
        report[
            "decision_counts"
        ].items()
    ):
        print(
            f"  {key}: {value}"
        )

    print()

    print(
        "role_counts:"
    )

    for key, value in (
        report[
            "role_counts"
        ].items()
    ):
        print(
            f"  {key}: {value}"
        )

    print()

    print(
        "category_counts:"
    )

    for key, value in (
        report[
            "category_counts"
        ].items()
    ):
        print(
            f"  {key}: {value}"
        )

    print()

    print(
        "clarify_case_count: "
        f"{report['clarify_case_count']}"
    )

    print()
    print(
        "Structured Gold intentionally "
        "left unavailable:"
    )

    for field in (
        report[
            "unavailable_structured_gold"
        ]
    ):
        print(
            f"  - {field}"
        )

    print()

    print(
        "conversion_report: "
        f"{REPORT_JSON}"
    )

    print()
    print(
        "VALIDATION: PASSED"
    )

    print("=" * 100)


# ================================================================
# Main
# ================================================================
def main():

    print(
        f"Loading frozen dataset: "
        f"{SOURCE_XLSX}"
    )

    source_cases = (
        load_eval_cases()
    )

    print(
        f"Loaded "
        f"{len(source_cases)} "
        f"cases."
    )

    print(
        "Validating frozen source..."
    )

    validate_source_cases(
        source_cases
    )

    print(
        "Frozen source validation passed."
    )

    converted_cases = [
        convert_case(
            source_case
        )
        for source_case
        in source_cases
    ]

    print(
        "Validating converted Core 61..."
    )

    report = (
        validate_converted_cases(
            converted_cases
        )
    )

    print(
        "Converted Core 61 validation passed."
    )

    write_csv(
        converted_cases
    )

    write_report(
        report
    )

    print_summary(
        report
    )


if __name__ == "__main__":
    main()