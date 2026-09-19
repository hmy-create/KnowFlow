from pathlib import Path

from docx import Document


REPO_ROOT = Path(__file__).resolve().parents[1]

CORPUS_DIR = (
    REPO_ROOT
    / "docs"
    / "corpus"
    / "KnowFlow_Corpus_V1.0"
)

OUTPUT_FILE = (
    REPO_ROOT
    / "docs"
    / "corpus"
    / "corpus_docs_preview.txt"
)


def clean_text(text: str) -> str:
    return " ".join(
        text.replace("\n", " ").split()
    )


def main():
    lines = []

    files = sorted(
        CORPUS_DIR.glob("*.docx")
    )

    for path in files:

        # 00 是说明文档，前面已经单独检查
        if path.name.startswith("00_"):
            continue

        lines.append("")
        lines.append(
            "=" * 80
        )
        lines.append(
            f"FILE: {path.name}"
        )
        lines.append(
            "=" * 80
        )

        doc = Document(path)

        # 先输出前 20 个非空段落
        paragraph_count = 0

        for paragraph in doc.paragraphs:
            text = clean_text(
                paragraph.text
            )

            if not text:
                continue

            lines.append(
                f"[P] {text}"
            )

            paragraph_count += 1

            if paragraph_count >= 20:
                break

        # 表格通常可能保存版本元数据
        for table_index, table in enumerate(
            doc.tables,
            start=1,
        ):
            lines.append(
                f"[TABLE {table_index}]"
            )

            for row in table.rows[:10]:
                cells = [
                    clean_text(cell.text)
                    for cell in row.cells
                ]

                lines.append(
                    " | ".join(cells)
                )

    OUTPUT_FILE.write_text(
        "\n".join(lines),
        encoding="utf-8",
    )

    print(
        "Corpus preview generated:"
    )
    print(
        OUTPUT_FILE
    )


if __name__ == "__main__":
    main()